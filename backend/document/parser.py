import re
from pathlib import Path
from zipfile import BadZipFile

from docx import Document
from docx.oxml.ns import qn

from .models import Citation, DocumentAnalysis, Equation, Figure, Reference, Section, TableData

SECTION_ALIASES = {
    "introduction": "introduction",
    "background": "background",
    "related work": "related_work",
    "literature review": "related_work",
    "methodology": "methodology",
    "materials and methods": "methodology",
    "materials & methods": "methodology",
    "proposed method": "methodology",
    "proposed methods": "methodology",
    "methods": "methodology",
    "results": "results",
    "discussion": "discussion",
    "conclusion": "conclusion",
    "future work": "future_work",
    "limitations": "limitations",
    "references": "references",
    "bibliography": "references",
    "abstract": "abstract",
    "keywords": "keywords",
}
CITATION_PATTERNS = [
    ("numeric", re.compile(r"\[(\d+(?:\s*[-,]\s*\d+)*)\]")),
    ("author_year", re.compile(r"\b([A-Z][A-Za-z-]+(?:\s+et al\.)?)\s*\((\d{4}[a-z]?)\)")),
]

class DocumentParser:
    def parse(self, path: Path, document_id: str) -> DocumentAnalysis:
        try:
            document = Document(path)
        except (BadZipFile, ValueError, OSError) as exc:
            raise ValueError("The uploaded file is not a readable DOCX document.") from exc
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        sections = self._sections(document)
        references = self._references(sections)
        return DocumentAnalysis(
            document_id=document_id,
            title=self._title(paragraphs, sections),
            authors=self._authors(paragraphs),
            affiliations=self._affiliations(paragraphs),
            abstract=self._abstract(sections),
            keywords=self._keywords(paragraphs),
            paragraphs=paragraphs,
            sections=sections,
            figures=self._figures(document, sections),
            tables=self._tables(document, sections),
            equations=self._equations(document, sections),
            citations=self._citations(paragraphs, sections),
            references=references,
            metadata={"paragraph_count": len(paragraphs), "table_count": len(document.tables), "inline_shape_count": len(document.inline_shapes)},
        )

    def _sections(self, document: Document) -> list[Section]:
        sections: list[Section] = []
        current: Section | None = None
        position = 0
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name.lower() if paragraph.style else ""
            is_heading = self._is_heading(text, style)
            if is_heading:
                if current:
                    sections.append(current)
                level = self._heading_level(text, style)
                normalized = self._normalize_heading(text)
                position += 1
                current = Section(section_id=f"SEC-{position:02d}", heading=text, type=normalized, level=level, position=position)
            elif current:
                current.content = f"{current.content}\n{text}".strip()
        if current:
            sections.append(current)
        if not sections and document.paragraphs:
            content = "\n".join(p.text.strip() for p in document.paragraphs if p.text.strip())
            sections.append(Section(section_id="SEC-01", heading="Body", type="other", content=content, position=1))
        return sections

    def _is_heading(self, text: str, style: str) -> bool:
        if not text:
            return False
        clean = text.strip()
        if style.startswith("heading"):
            return True
        if re.match(r"^(?:\d+|[IVXLC]+)(?:[.)]|\.|\s+)", clean):
            return len(clean.split()) <= 12
        if re.match(r"^(?:[A-Z][A-Za-z0-9()\-/& ]{2,80})$", clean) and len(clean.split()) <= 12:
            if any(alias in clean.lower() for alias in SECTION_ALIASES):
                return True
        return bool(re.match(r"^(?:\d+(?:\.\d+)*[.)]?\s+)?[A-Z][^.!?]{2,80}$", clean)) and len(clean.split()) <= 12

    def _heading_level(self, text: str, style: str) -> int:
        if style.startswith("heading"):
            match = re.search(r"(\d+)", style)
            return int(match.group(1)) if match else 1
        if re.match(r"^(?:\d+)(?:\.\d+)*[.)]?\s+", text):
            numbering = re.match(r"^(?:\d+)(?:\.\d+)*", text)
            if numbering:
                return max(1, len(numbering.group(0).split(".")))
        return 1

    def _normalize_heading(self, heading: str) -> str:
        clean = re.sub(r"^\d+(?:\.\d+)*[.)]?\s+", "", heading.lower()).strip(" :")
        clean = clean.strip()
        if not clean:
            return "other"
        return SECTION_ALIASES.get(clean, "other")

    def _title(self, paragraphs: list[str], sections: list[Section]) -> str:
        for heading in (section.heading for section in sections if section.type not in {"abstract", "references", "keywords"}):
            if len(heading.split()) <= 20:
                return heading
        for paragraph in paragraphs[:10]:
            if not any(token in paragraph.lower() for token in ("abstract", "keywords", "university", "department")) and len(paragraph.split()) <= 30:
                return paragraph
        return sections[0].heading if sections else "Untitled manuscript"

    def _authors(self, paragraphs: list[str]) -> list[str]:
        for paragraph in paragraphs[1:5]:
            if "," in paragraph and len(paragraph.split()) < 30 and not any(word in paragraph.lower() for word in ("university", "department", "institute")):
                return [item.strip() for item in paragraph.split(",") if item.strip()]
        return []

    def _affiliations(self, paragraphs: list[str]) -> list[str]:
        return [paragraph for paragraph in paragraphs[:8] if any(word in paragraph.lower() for word in ("university", "department", "institute", "laboratory"))]

    def _abstract(self, sections: list[Section]) -> str:
        return next((section.content for section in sections if section.type == "abstract"), "")

    def _keywords(self, paragraphs: list[str]) -> list[str]:
        for paragraph in paragraphs:
            if paragraph.lower().startswith(("keywords:", "key words:")):
                return [item.strip() for item in re.split(r"[:,;]", paragraph, maxsplit=1)[-1].split(",") if item.strip()]
        return []

    def _figures(self, document: Document, sections: list[Section]) -> list[Figure]:
        figures = []
        for index, shape in enumerate(document.inline_shapes, start=1):
            caption = ""
            if shape._inline.xpath(".//wp:docPr"):
                name = shape._inline.xpath(".//wp:docPr")[0].get("descr") or ""
                caption = name
            figures.append(Figure(id=f"Figure {index}", caption=caption, position=index, image_name=f"figure_{index}.png"))
        return figures

    def _tables(self, document: Document, sections: list[Section]) -> list[TableData]:
        tables = []
        for index, table in enumerate(document.tables, start=1):
            content = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            tables.append(TableData(id=f"Table {index}", content=content, position=index))
        return tables

    def _equations(self, document: Document, sections: list[Section]) -> list[Equation]:
        equations = []
        for index, element in enumerate(document.element.body.iter(), start=1):
            if element.tag == qn("m:oMath") or element.tag == qn("m:oMathPara"):
                equations.append(Equation(id=f"Equation {len(equations) + 1}", position=index, xml=element.xml))
        return equations

    def _citations(self, paragraphs: list[str], sections: list[Section]) -> list[Citation]:
        citations = []
        position = 0
        for paragraph_index, paragraph in enumerate(paragraphs):
            for citation_type, pattern in CITATION_PATTERNS:
                for match in pattern.finditer(paragraph):
                    position += 1
                    reference_ids = self._expand_numeric(match.group(1)) if citation_type == "numeric" else []
                    citations.append(Citation(text=match.group(0), citation_type=citation_type, location=f"paragraph:{paragraph_index + 1}", position=position, reference_ids=reference_ids))
        return citations

    def _expand_numeric(self, raw: str) -> list[str]:
        values = []
        for token in re.split(r"\s*,\s*", raw):
            if "-" in token:
                start, end = [int(value.strip()) for value in token.split("-", 1)]
                values.extend(str(value) for value in range(start, end + 1))
            elif token.strip().isdigit():
                values.append(token.strip())
        return values

    def _references(self, sections: list[Section]) -> list[Reference]:
        reference_section = next((section for section in sections if section.type == "references"), None)
        if not reference_section:
            return []
        entries = [line.strip() for line in reference_section.content.splitlines() if line.strip()]
        references = []
        for index, raw in enumerate(entries, start=1):
            number_match = re.match(r"\[?(\d+)\]?\.?\s+(.+)", raw)
            reference_id = number_match.group(1) if number_match else str(index)
            text = number_match.group(2) if number_match else raw
            year = (re.search(r"\b(19|20)\d{2}[a-z]?\b", text) or [None])[0]
            doi_match = re.search(r"10\.\d{4,9}/\S+", text, re.I)
            url_match = re.search(r"https?://\S+", text)
            references.append(Reference(id=reference_id, raw_text=raw, year=year, doi=doi_match.group(0).rstrip(".,") if doi_match else None, url=url_match.group(0).rstrip(".,") if url_match else None))
        return references
