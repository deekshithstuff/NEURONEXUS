from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


PAGE_SIZES = {
    "A4": (Inches(8.27), Inches(11.69)),
    "LETTER": (Inches(8.5), Inches(11.0)),
}


def _set_run_font(run, font_name: str, font_size: int, bold: bool = False) -> None:
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), font_name)
    rfonts.set(qn("w:hAnsi"), font_name)
    rfonts.set(qn("w:eastAsia"), font_name)


class DOCXGenerator:
    def generate(self, document: dict, output_path: str | Path, journal_rules: dict | None = None) -> str:
        rules = journal_rules or {}
        font_name = rules.get("font") or "Times New Roman"
        font_size = int(rules.get("font_size") or 11)
        line_spacing = float(rules.get("line_spacing") or 1.15)
        paragraph_spacing = int(rules.get("paragraph_spacing") or 6)
        margins = rules.get("margins") or {"top": 1.0, "bottom": 1.0, "left": 1.0, "right": 1.0}

        doc = Document()
        section = doc.sections[0]
        page_key = str(rules.get("page_size") or "A4").upper()
        width, height = PAGE_SIZES.get(page_key, PAGE_SIZES["A4"])
        section.page_width = width
        section.page_height = height
        section.top_margin = Inches(float(margins.get("top", 1.0)))
        section.bottom_margin = Inches(float(margins.get("bottom", 1.0)))
        section.left_margin = Inches(float(margins.get("left", 1.0)))
        section.right_margin = Inches(float(margins.get("right", 1.0)))

        title = document.get("title") or "Untitled manuscript"
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_para.add_run(title)
        _set_run_font(title_run, font_name, font_size + 4, bold=True)

        authors = document.get("authors") or []
        if authors:
            author_para = doc.add_paragraph()
            author_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            author_run = author_para.add_run(", ".join(authors))
            _set_run_font(author_run, font_name, font_size)

        keywords = document.get("keywords") or []
        if keywords:
            keyword_para = doc.add_paragraph()
            keyword_run = keyword_para.add_run("Keywords: " + ", ".join(str(item) for item in keywords))
            _set_run_font(keyword_run, font_name, font_size, bold=False)
            self._apply_spacing(keyword_para, line_spacing, paragraph_spacing)

        for manuscript_section in document.get("sections", []):
            heading = manuscript_section.get("heading")
            if heading:
                heading_para = doc.add_heading(heading, level=min(int(manuscript_section.get("level") or 1), 4))
                for run in heading_para.runs:
                    _set_run_font(run, font_name, font_size + 2, bold=True)
            content = manuscript_section.get("content")
            if content:
                body = doc.add_paragraph(content)
                self._apply_spacing(body, line_spacing, paragraph_spacing)
                for run in body.runs:
                    _set_run_font(run, font_name, font_size)

        for index, figure in enumerate(document.get("figures") or [], start=1):
            caption = figure.get("caption") or figure.get("id") or f"Figure {index}"
            para = doc.add_paragraph(f"Figure {index}. {caption}")
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                _set_run_font(run, font_name, max(font_size - 1, 8))

        for index, table in enumerate(document.get("tables") or [], start=1):
            caption = table.get("caption") or table.get("id") or f"Table {index}"
            para = doc.add_paragraph(f"Table {index}. {caption}")
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                _set_run_font(run, font_name, max(font_size - 1, 8))
            rows = table.get("content") or table.get("rows") or []
            if isinstance(rows, list) and rows and isinstance(rows[0], list):
                word_table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                for r_idx, row in enumerate(rows):
                    for c_idx, cell_value in enumerate(row):
                        word_table.rows[r_idx].cells[c_idx].text = str(cell_value)

        for index, equation in enumerate(document.get("equations") or [], start=1):
            text = equation.get("latex") or equation.get("text") or equation.get("id") or f"Equation {index}"
            para = doc.add_paragraph(f"({index})  {text}")
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                _set_run_font(run, font_name, font_size)

        references = document.get("references") or []
        has_reference_section = any(
            str((section or {}).get("type") or "").lower() == "references"
            and str((section or {}).get("content") or "").strip()
            for section in document.get("sections") or []
        )
        if references and not has_reference_section:
            heading = doc.add_heading("References", level=1)
            for run in heading.runs:
                _set_run_font(run, font_name, font_size + 2, bold=True)
            for index, reference in enumerate(references, start=1):
                raw = reference.get("raw_text") or reference.get("text") or ""
                ref_id = reference.get("id") or str(index)
                para = doc.add_paragraph(f"[{ref_id}] {raw}".strip())
                self._apply_spacing(para, line_spacing, paragraph_spacing)
                for run in para.runs:
                    _set_run_font(run, font_name, font_size)

        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return str(output_path)

    def _apply_spacing(self, paragraph, line_spacing: float, paragraph_spacing: int) -> None:
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        paragraph.paragraph_format.line_spacing = line_spacing
        paragraph.paragraph_format.space_after = Pt(paragraph_spacing)
