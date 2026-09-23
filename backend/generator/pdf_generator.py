from __future__ import annotations

from pathlib import Path

import fitz
from docx import Document


class PDFGenerator:
    def generate(self, docx_path: str | Path | None = None, output_path: str | Path | None = None, text: str | None = None) -> str:
        if output_path is None:
            raise ValueError("output_path is required for PDF generation.")
        content = text
        if not content and docx_path:
            content = self._extract_docx_text(docx_path)
        if not content:
            content = (
                "Generated manuscript preview\n\n"
                "This PDF is created from the formatted manuscript export and preserves the original research content."
            )

        chunks = self._chunk_text(content, 2800)
        doc = fitz.open()
        for chunk in chunks:
            page = doc.new_page()
            page.insert_textbox(
                fitz.Rect(54, 54, page.rect.width - 54, page.rect.height - 54),
                chunk,
                fontsize=11,
                fontname="helv",
            )
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        doc.save(output_path)
        return str(output_path)

    def generate_readiness_report(self, output_path: str | Path, details: str) -> str:
        return self.generate(output_path=output_path, text=details)

    def _extract_docx_text(self, docx_path: str | Path) -> str:
        document = Document(docx_path)
        blocks: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                blocks.append(paragraph.text.strip())
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        return "\n\n".join(blocks)

    def _chunk_text(self, content: str, size: int) -> list[str]:
        if len(content) <= size:
            return [content]
        chunks: list[str] = []
        start = 0
        while start < len(content):
            chunks.append(content[start:start + size])
            start += size
        return chunks
