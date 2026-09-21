from __future__ import annotations

from pathlib import Path

from fastapi.testclient import TestClient
from docx import Document

from backend.main import app

client = TestClient(app)


def build_sample_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Deep Learning for Clinical Decision Support", level=1)
    doc.add_paragraph("Alice Smith, Bob Jones")
    doc.add_paragraph("Department of Computer Science, University of Example")
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("This study evaluates a clinical decision support model. We present a method and results. [1], [2], [3]")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Recent work shows value in expert systems. Smith et al. (2024) and (Brown, 2023) highlight prior findings.")
    doc.add_heading("Methods", level=1)
    doc.add_paragraph("We trained a convolutional model on clinical records.")
    doc.add_heading("Results", level=1)
    doc.add_paragraph("The model achieved strong performance.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Smith A. et al. Deep learning in medicine. Nature 2024.")
    doc.add_paragraph("[2] Brown K. Clinical systems. IEEE TMI 2023.")
    doc.add_paragraph("[3] Jones R. Decision support views. AI Med 2025.")
    doc.save(path)


def test_upload_and_analysis_flow(tmp_path):
    sample_path = tmp_path / "sample.docx"
    build_sample_docx(sample_path)

    with sample_path.open("rb") as fh:
        response = client.post("/api/documents/upload", files={"file": ("sample.docx", fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200, response.text
    payload = response.json()
    doc_id = payload["document_id"]
    assert payload["status"] == "uploaded"

    analysis = client.post(f"/api/documents/{doc_id}/analyze", json={"journal_id": "nature"})
    assert analysis.status_code == 200, analysis.text
    body = analysis.json()
    assert body["document_id"] == doc_id
    assert body["title"]
    assert len(body["sections"]) >= 3
    assert body["citations"]
    assert body["references"]

    citations_report = client.get(f"/api/citations/{doc_id}")
    assert citations_report.status_code == 200, citations_report.text
    report = citations_report.json()
    assert report["total_citations"] >= 1
    assert report["total_references"] >= 1


def test_author_year_citations_in_numbered_sections(tmp_path):
    sample_path = tmp_path / "author_year.docx"
    doc = Document()
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("Smith and Johnson (2024) report a strong benchmark. Davis et al. (2023) confirms the trend.")
    doc.add_heading("2. Methods", level=1)
    doc.add_paragraph("We compare against the prior work.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("1. Smith, J. and Johnson, K. Clinical benchmarking. Journal of Methods. 2024.")
    doc.add_paragraph("2. Davis, L. et al. Prior trends in evaluation. ML Review. 2023.")
    doc.save(sample_path)

    with sample_path.open("rb") as fh:
        response = client.post("/api/documents/upload", files={"file": ("author_year.docx", fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert response.status_code == 200, response.text

    doc_id = response.json()["document_id"]
    analysis = client.post(f"/api/documents/{doc_id}/analyze", json={"journal_id": "nature"})
    assert analysis.status_code == 200, analysis.text

    payload = analysis.json()
    assert payload["title"]
    assert any(section["type"] == "introduction" for section in payload["sections"])
    assert len(payload["citations"]) >= 2
    assert any(citation["citation_type"] == "author_year" for citation in payload["citations"])


def test_author_year_citations_match_references():
    citations = [
        {"text": "Smith et al. (2024)", "citation_type": "author_year", "reference_ids": []},
        {"text": "Brown, 2023", "citation_type": "author_year", "reference_ids": []},
    ]
    references = [
        {"id": "1", "raw_text": "Smith J. et al. Clinical benchmarking. Journal of Methods. 2024.", "authors": ["Smith J."], "year": "2024"},
        {"id": "2", "raw_text": "Brown K. Clinical systems. IEEE TMI 2023.", "authors": ["Brown K."], "year": "2023"},
    ]

    report = client.post("/api/citations/check", json={"citations": citations, "references": references})
    assert report.status_code == 200, report.text
    body = report.json()
    assert body["missing_references"] == []
    assert body["uncited_references"] == []


def test_journal_rules_and_generation_flow(tmp_path):
    sample_path = tmp_path / "journal_sample.docx"
    build_sample_docx(sample_path)

    with sample_path.open("rb") as fh:
        upload = client.post("/api/documents/upload", files={"file": ("journal_sample.docx", fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    doc_id = upload.json()["document_id"]
    client.post(f"/api/documents/{doc_id}/analyze", json={"journal_id": "nature"})

    journals = client.get("/api/journals")
    assert journals.status_code == 200
    assert "journals" in journals.json()

    rules = client.get("/api/journals/nature/rules")
    assert rules.status_code == 200
    assert rules.json()["rules"]["journal_id"] == "nature"

    format_response = client.post(f"/api/documents/{doc_id}/format", json={"journal_id": "nature"})
    assert format_response.status_code == 200, format_response.text
    assert format_response.json()["journal_id"] == "nature"

    generate = client.post(f"/api/documents/{doc_id}/generate")
    assert generate.status_code == 200, generate.text
    assert generate.json()["status"] == "generated"

    submission = generate.json()["submission_package"]
    final_docx_path = Path(submission["final_docx"])
    assert final_docx_path.exists()
    assert Path(submission["final_pdf"]).exists()
    assert Path(submission["readiness_report"]).exists()

    generated_doc = Document(final_docx_path)
    texts = [p.text for p in generated_doc.paragraphs if p.text.strip()]
    assert any("Deep Learning for Clinical Decision Support" in text for text in texts)
    assert any("This study evaluates a clinical decision support model." in text for text in texts)
    assert any("References" in text for text in texts)

    docx_download = client.get(f"/api/documents/{doc_id}/download/docx")
    assert docx_download.status_code == 200

    pdf_download = client.get(f"/api/documents/{doc_id}/download/pdf")
    assert pdf_download.status_code == 200


def test_ai_quality_and_report_endpoints(tmp_path):
    sample_path = tmp_path / "ai_quality_sample.docx"
    doc = Document()
    doc.add_heading("Deep Learning for Clinical Diagnosis", level=1)
    doc.add_paragraph("A. Lee, B. Patel")
    doc.add_paragraph("Abstract: We propose a deep learning approach for diagnosis using chest imaging. The method uses CNNs and compares to baseline models. [1], [2]")
    doc.add_heading("Introduction", level=1)
    doc.add_paragraph("Existing work relies on handcrafted features. This gap motivates our CNN-based pipeline.")
    doc.add_heading("Methodology", level=1)
    doc.add_paragraph("We use a ResNet-50 model trained on 12,000 CT scans using Adam with learning rate 0.001.")
    doc.add_heading("Results", level=1)
    doc.add_paragraph("The model achieved Macro-F1 of 0.89 and AUROC of 0.94.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Smith J. Medical imaging AI. Nature 2024.")
    doc.add_paragraph("[2] Jones A. Clinical computer vision. IEEE 2023.")
    doc.save(sample_path)

    with sample_path.open("rb") as fh:
        upload = client.post("/api/documents/upload", files={"file": ("ai_quality_sample.docx", fh, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    doc_id = upload.json()["document_id"]
    analysis = client.post(f"/api/documents/{doc_id}/analyze", json={"journal_id": "nature"})
    assert analysis.status_code == 200, analysis.text

    quality = client.post("/api/quality/analyze", json={"document_id": doc_id, "analysis": analysis.json()})
    assert quality.status_code == 200, quality.text
    quality_body = quality.json()
    assert quality_body["research_gap"]
    assert "missing_sections" in quality_body

    match = client.post("/api/journal/match", json={"document_id": doc_id, "analysis": analysis.json(), "journal_id": "nature"})
    assert match.status_code == 200, match.text
    assert match.json()["journal_id"] == "nature"

    report = client.post("/api/report/generate", json={"document_id": doc_id, "analysis": analysis.json(), "quality_analysis": quality_body, "selected_journal": "nature"})
    assert report.status_code == 200, report.text
    assert "Pre-Submission Readiness Assessment" in report.json()["title"]


def test_journal_match_uses_real_manuscript_content_and_journal_scope():
    manuscript_a = {
        "title": "IoT-Based Smart Irrigation System Using Soil Moisture Monitoring and Machine Learning",
        "abstract": "This paper presents an IoT-based smart irrigation system for agriculture using soil moisture sensing, ESP32 controllers, and machine learning models for adaptive watering control.",
        "keywords": ["IoT", "Smart Irrigation", "Soil Moisture", "ESP32", "Machine Learning", "Agriculture"],
        "paragraphs": [
            "The goal is to monitor soil moisture in agricultural fields using ESP32 sensors and machine learning-driven irrigation control.",
            "We evaluate irrigation strategies for water efficiency and crop health in rural agriculture settings.",
        ],
        "sections": [
            {"type": "abstract", "heading": "Abstract", "content": "IoT-based smart irrigation..."},
            {"type": "methodology", "heading": "Methodology", "content": "ESP32 devices collect soil moisture from agricultural plots."},
            {"type": "results", "heading": "Results", "content": "Machine learning improved irrigation decisions in the agricultural field experiment."},
        ],
    }

    manuscript_b = {
        "title": "Satellite-Based Crop Disease Detection Using Computer Vision",
        "abstract": "We propose a computer vision pipeline for detecting crop disease from high-resolution satellite imagery and agricultural field observations.",
        "keywords": ["satellite imagery", "crop disease", "computer vision", "deep learning", "agriculture"],
        "paragraphs": [
            "Our method uses satellite imagery to identify disease symptoms across agricultural regions using computer vision models.",
            "The pipeline combines remote sensing and deep learning for crop health monitoring.",
        ],
        "sections": [
            {"type": "abstract", "heading": "Abstract", "content": "Computer vision detects crop disease from satellite imagery."},
            {"type": "methodology", "heading": "Methodology", "content": "Remote sensing and deep learning support crop disease classification in agriculture."},
            {"type": "results", "heading": "Results", "content": "The computer vision approach shows strong crop disease detection performance."},
        ],
    }

    result_a_nature = client.post("/api/journal/match", json={"document_id": "DOC-A", "analysis": manuscript_a, "journal_id": "nature"})
    result_b_nature = client.post("/api/journal/match", json={"document_id": "DOC-B", "analysis": manuscript_b, "journal_id": "nature"})
    result_a_ieee = client.post("/api/journal/match", json={"document_id": "DOC-A", "analysis": manuscript_a, "journal_id": "ieee"})
    result_b_ieee = client.post("/api/journal/match", json={"document_id": "DOC-B", "analysis": manuscript_b, "journal_id": "ieee"})

    assert result_a_nature.status_code == 200
    assert result_b_nature.status_code == 200
    assert result_a_ieee.status_code == 200
    assert result_b_ieee.status_code == 200

    a_nature = result_a_nature.json()
    b_nature = result_b_nature.json()
    a_ieee = result_a_ieee.json()
    b_ieee = result_b_ieee.json()

    assert a_nature["journal_id"] == "nature"
    assert a_ieee["journal_id"] == "ieee"
    assert b_nature["journal_id"] == "nature"
    assert b_ieee["journal_id"] == "ieee"

    assert a_nature["relevant_topics"] != b_nature["relevant_topics"]
    assert a_ieee["relevant_topics"] != b_ieee["relevant_topics"]
    assert a_nature["relevant_topics"] != ["medical imaging", "clinical AI", "deep learning", "evaluation"]
    assert a_ieee["relevant_topics"] != ["medical imaging", "clinical AI", "deep learning", "evaluation"]
    assert a_nature["explanation"] != b_nature["explanation"]
    assert a_ieee["explanation"] != b_ieee["explanation"]
    assert 0.0 < a_nature["score"] < 1.0
    assert 0.0 < a_ieee["score"] < 1.0
    assert 0.0 < b_nature["score"] < 1.0
    assert 0.0 < b_ieee["score"] < 1.0

    assert set(a_ieee["relevant_topics"]).intersection({"iot", "soil moisture", "smart irrigation", "esp32", "agriculture"})
    assert set(b_ieee["relevant_topics"]).intersection({"satellite imagery", "crop disease", "computer vision", "agriculture"})

    assert "suitable" in a_nature and "suitable" in a_ieee
    assert isinstance(a_nature["suitable"], bool)
    assert isinstance(a_ieee["suitable"], bool)

    missing = {"title": "", "abstract": "", "keywords": []}
    empty_match = client.post("/api/journal/match", json={"document_id": "DOC-EMPTY", "analysis": missing, "journal_id": "nature"})
    assert empty_match.status_code == 200
    empty_body = empty_match.json()
    assert empty_body["journal_id"] == "nature"
    assert 0.0 <= empty_body["score"] <= 1.0
    assert isinstance(empty_body["relevant_topics"], list)
    assert isinstance(empty_body["scope_gaps"], list)
