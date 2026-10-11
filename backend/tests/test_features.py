from __future__ import annotations

from io import BytesIO
from pathlib import Path
from uuid import uuid4

from docx import Document
from fastapi.testclient import TestClient

from backend.main import app

client = TestClient(app)


def new_client() -> tuple[TestClient, dict]:
    api = TestClient(app)
    response = api.post(
        "/api/auth/register",
        json={
            "name": f"Feature user {uuid4()}",
            "email": f"feature-{uuid4()}@example.test",
            "password": "correct-horse-battery",
        },
    )
    assert response.status_code == 200, response.text
    account = response.json()
    api.headers.update({"Authorization": f"Bearer {account['access_token']}"})
    return api, account


def build_manuscript() -> BytesIO:
    document = Document()
    document.add_heading("Statistical Models for Renewable Energy Forecasting", level=1)
    document.add_paragraph("Dana Lee, Chris Ford")
    document.add_heading("Abstract", level=1)
    document.add_paragraph("We present a forecasting model and report results. [1], [2]")
    document.add_heading("Introduction", level=1)
    document.add_paragraph("Prior work is limited; there is a clear research gap in this area. Lee et al. (2024).")
    document.add_heading("Methods", level=1)
    document.add_paragraph("We trained a regression model on hourly data and evaluated it on a held-out test set.")
    document.add_heading("Results", level=1)
    document.add_paragraph("The model reduced error by 12% over the baseline.")
    document.add_heading("References", level=1)
    document.add_paragraph("[1] Lee D. Forecasting models. Energy Journal 2024.")
    document.add_paragraph("[2] Ford C. Statistical evaluation. Applied Energy 2023.")
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def analyze_document(api: TestClient, journal_id: str = "nature") -> str:
    upload = api.post(
        "/api/documents/upload",
        files={"file": ("manuscript.docx", build_manuscript(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert upload.status_code == 200, upload.text
    document_id = upload.json()["document_id"]
    analyzed = api.post(f"/api/documents/{document_id}/analyze", json={"journal_id": journal_id})
    assert analyzed.status_code == 200, analyzed.text
    return document_id


def test_compliance_check_reports_rule_results():
    api, _ = new_client()
    document_id = analyze_document(api)
    response = api.get(f"/api/documents/{document_id}/compliance")
    assert response.status_code == 200, response.text
    report = response.json()
    assert report["document_id"] == document_id
    assert report["overall_status"] in {"passed", "warning", "failed"}
    assert report["rules"], "compliance report should contain rule results"
    assert all({"rule_id", "status", "category"} <= set(rule) for rule in report["rules"])


def test_compliance_requires_access():
    owner_api, _ = new_client()
    document_id = analyze_document(owner_api)
    outsider_api, _ = new_client()
    response = outsider_api.get(f"/api/documents/{document_id}/compliance")
    assert response.status_code == 404


def test_analytics_dashboard_reports_indicators():
    api, _ = new_client()
    document_id = analyze_document(api)
    response = api.get(f"/api/documents/{document_id}/analytics")
    assert response.status_code == 200, response.text
    dashboard = response.json()
    keys = {indicator["key"] for indicator in dashboard["indicators"]}
    assert {"submission_readiness", "compliance", "completeness", "improvement", "similarity"} <= keys
    readiness = next(item for item in dashboard["indicators"] if item["key"] == "submission_readiness")
    assert readiness["available"] is True
    assert readiness["value"] is not None
    similarity = next(item for item in dashboard["indicators"] if item["key"] == "similarity")
    assert similarity["available"] is False
    assert "does not prove plagiarism" in dashboard["disclaimer"] or "do not detect plagiarism" in dashboard["disclaimer"]
    assert dashboard["snapshot_id"]

    exported = api.get(f"/api/documents/{document_id}/analytics/export")
    assert exported.status_code == 200
    assert "attachment" in exported.headers.get("content-disposition", "")


def test_conversion_creates_trackable_jobs():
    api, _ = new_client()
    document_id = analyze_document(api)
    response = api.post(
        f"/api/documents/{document_id}/conversions",
        json={"journal_ids": ["ieee", "nature"], "article_type": "research_article"},
    )
    assert response.status_code == 200, response.text
    jobs = response.json()["jobs"]
    assert {job["journal_id"] for job in jobs} == {"ieee", "nature"}

    listed = api.get(f"/api/documents/{document_id}/conversions")
    assert listed.status_code == 200
    assert listed.json()["jobs"]

    for job in jobs:
        detail = api.get(f"/api/conversions/{job['job_id']}")
        assert detail.status_code == 200, detail.text
        assert detail.json()["status"] in {"queued", "running", "completed", "failed"}


def test_conversion_rejects_unknown_journal():
    api, _ = new_client()
    document_id = analyze_document(api)
    response = api.post(
        f"/api/documents/{document_id}/conversions",
        json={"journal_ids": ["not-a-real-journal"]},
    )
    assert response.status_code in {400, 404}


def test_collaboration_invite_comment_review_and_versions():
    owner_api, owner = new_client()
    document_id = analyze_document(owner_api)

    reviewer_api, reviewer = new_client()
    invite = owner_api.post(
        f"/api/documents/{document_id}/collaborators",
        json={"email": reviewer["user"]["email"], "role": "reviewer"},
    )
    assert invite.status_code == 200, invite.text
    assert invite.json()["status"] == "accepted"

    shared = reviewer_api.get(f"/api/documents/{document_id}")
    assert shared.status_code == 200, shared.text

    comment = reviewer_api.post(
        f"/api/documents/{document_id}/comments",
        json={"body": "Please clarify the evaluation protocol.", "section_ref": "Methods"},
    )
    assert comment.status_code == 200, comment.text
    comment_id = comment.json()["comment_id"] if "comment_id" in comment.json() else comment.json().get("id")
    assert comment_id

    comments = reviewer_api.get(f"/api/documents/{document_id}/comments")
    assert comments.status_code == 200

    version = owner_api.post(
        f"/api/documents/{document_id}/versions",
        json={"label": "v1", "summary": "Initial revision"},
    )
    assert version.status_code == 200, version.text
    versions = owner_api.get(f"/api/documents/{document_id}/versions")
    assert versions.status_code == 200

    review = reviewer_api.post(
        f"/api/documents/{document_id}/reviews",
        json={"decision": "revision_requested", "comment": "Solid contribution; clarify methods."},
    )
    assert review.status_code == 200, review.text
    decisions = reviewer_api.get(f"/api/documents/{document_id}/reviews")
    assert decisions.status_code == 200
    assert decisions.json()["latest"]["decision"] == "revision_requested"

    forbidden = reviewer_api.post(
        f"/api/documents/{document_id}/collaborators",
        json={"email": f"nobody-{uuid4()}@example.test", "role": "viewer"},
    )
    assert forbidden.status_code == 403


def test_viewer_cannot_comment():
    owner_api, _ = new_client()
    document_id = analyze_document(owner_api)
    viewer_api, viewer = new_client()
    invite = owner_api.post(
        f"/api/documents/{document_id}/collaborators",
        json={"email": viewer["user"]["email"], "role": "viewer"},
    )
    assert invite.status_code == 200, invite.text
    response = viewer_api.post(
        f"/api/documents/{document_id}/comments",
        json={"body": "Should not be allowed."},
    )
    assert response.status_code == 403


def test_improvement_assessment_and_revision_are_advisory():
    api, _ = new_client()
    document_id = analyze_document(api)
    assessment = api.post("/api/improvement/assess", json={"document_id": document_id})
    assert assessment.status_code == 200, assessment.text
    payload = assessment.json()
    assert payload["recommendations"]
    assert all(item["requires_approval"] is True for item in payload["recommendations"])
    assert all({"evidence", "evidence_type", "confidence", "priority"} <= set(item) for item in payload["recommendations"])

    revision = api.post(
        "/api/improvement/revise",
        json={"document_id": document_id, "targets": [{"original": "This is a very very long sentence that should be reduced.", "goal": "concise"}]},
    )
    assert revision.status_code == 200, revision.text
    body = revision.json()
    assert body["source_modified"] is False
    assert body["requires_approval"] is True

    experiments = api.post("/api/improvement/experiments", json={"document_id": document_id})
    assert experiments.status_code == 200
    assert "disclaimer" in experiments.json()
